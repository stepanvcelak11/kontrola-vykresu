"""Testy importu tabulky, pravidel YAML, projektu (.kontrola), vzoru a exportů."""

import csv
import zipfile
from pathlib import Path

import ezdxf
import pytest
from conftest import check

from kontrola.config import Config
from kontrola.importer import template as tpl
from kontrola.importer.table import detect_header_row, generate_rules, guess_mapping, import_table
from kontrola.model import GeomType
from kontrola.project import Project
from kontrola.rules import Rule, RuleSet
from kontrola.runner import compare, run_checks

UKAZKY = Path(__file__).resolve().parents[1] / "ukazky"


def _csv(tmp_path, rows, name="t.csv", delim=";", enc="utf-8-sig"):
    p = tmp_path / name
    with open(p, "w", newline="", encoding=enc) as fh:
        csv.writer(fh, delimiter=delim).writerows(rows)
    return p


def test_odhad_hlavicky_a_sloupcu(tmp_path):
    p = _csv(tmp_path, [
        ["Cvičení z mapování – tabulka"],
        [],
        ["Poř.", "Kód", "Název objektu", "Level", "Barva", "Typ čáry", "Geometrie", "Povinné atributy",
         "Povolené hodnoty"],
        ["1", "101", "Budova", "BUDOVY", "červená", "plná", "plocha", "", ""],
        ["2", "402", "Strom", "VEGETACE", "3", "", "bod", "DRUH", "lípa, dub"],
        ["", "Komunikace", "", "", "", "", "", "", ""],  # nadpis skupiny – přeskočí se
        ["3", "", "Bez kódu", "X", "", "", "linie", "", ""],
        ["4", "101", "Duplicitní", "X", "", "", "", "", ""],
        ["5", "301", "Plot", "PLOTY", "5", "čárkovaná", "křivka?", "", ""],
    ], enc="cp1250", delim=",")
    td, hi, m = import_table(p)
    assert hi == 2
    assert m == {"kod": 1, "nazev": 2, "hladina": 3, "barva": 4, "styl_cary": 5, "geometrie": 6,
                 "povinne_atributy": 7, "povolene_hodnoty": 8}
    res = generate_rules(td.rows, hi, m, str(p))
    kody = [r.kod for r in res.rules.pravidla]
    assert kody == ["101", "402", "301"]
    assert len(res.errors) == 2  # bez kódu, duplicitní kód
    assert any("nerozpoznaný typ geometrie" in w for w in res.warnings)
    b = res.rules.by_code()
    assert b["101"].barva == "#FF0000" and b["101"].geometrie == GeomType.POLYGON  # „červená“
    assert b["101"].styl_cary == "CONTINUOUS"
    assert b["402"].povolene_hodnoty == {"DRUH": ["lípa", "dub"]}
    assert b["301"].styl_cary == "DASHED"


def test_import_ukazkove_tabulky_xlsx():
    td, hi, m = import_table(UKAZKY / "tabulka_atributu.xlsx")
    res = generate_rules(td.rows, hi, m)
    assert len(res.rules.pravidla) == 10
    assert res.errors == ["Řádek 14: neplatný kód „???“ (Lavička)."]


def test_import_pdf(tmp_path):
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table

    p = tmp_path / "tabulka.pdf"
    SimpleDocTemplate(str(p), pagesize=A4).build([Table([
        ["Kod", "Nazev", "Hladina", "Barva", "Typ"],
        ["101", "Budova", "BUDOVY", "1", "plocha"],
        ["201", "Silnice", "KOMUNIKACE", "8", "linie"],
    ], style=[("GRID", (0, 0), (-1, -1), 0.5, "black")])])
    td, hi, m = import_table(p)
    res = generate_rules(td.rows, hi, m)
    assert [r.kod for r in res.rules.pravidla] == ["101", "201"]


def test_pravidla_yaml_tam_a_zpet(tmp_path):
    rs = RuleSet(pravidla=[Rule(kod="402", nazev="Strom", geometrie=GeomType.BOD, blok="STROM",
                                povinne_atributy=["DRUH"], povolene_hodnoty={"DRUH": ["lípa"]})],
                 povolene_hladiny=["RAM"], rozsah="sjtsk")
    rs.save(tmp_path / "p.yaml")
    back = RuleSet.load(tmp_path / "p.yaml")
    assert back.to_dict() == rs.to_dict()


def test_ukazkova_konfigurace_se_nacte():
    rs = RuleSet.load(UKAZKY / "konfigurace.yaml")
    cfg = Config.load(UKAZKY / "konfigurace.yaml")
    assert len(rs.pravidla) >= 5
    assert cfg.tolerance == 0.05
    assert cfg.settings("visici_konce").zapnuto


def test_projekt_ulozeni_a_zip(tmp_path):
    p = Project.create(tmp_path / "projekt", "Úloha 3")
    p.set_drawing(UKAZKY / "ukazkovy_vykres.dxf")
    rel_img = p.add_attachment("obrazky", UKAZKY / "nacrt.png")
    p.set_image_note(rel_img, "bod 1005 chybí")
    p.rules = RuleSet(pravidla=[Rule(kod="1", hladina="A", obrazek=rel_img)])
    p.config.tolerance = 0.02
    p.save()
    out = p.export_zip(tmp_path / "moje.kontrola")
    assert out.suffix == ".kontrola"
    with zipfile.ZipFile(out) as z:
        names = set(z.namelist())
    assert {"projekt.yaml", "pravidla.yaml", "nastaveni.yaml", "vykres/ukazkovy_vykres.dxf",
            "podklady/obrazky/nacrt.png"} <= names
    q = Project.import_zip(out, tmp_path / "rozbaleno")
    assert q.name == "Úloha 3"
    assert q.config.tolerance == 0.02
    assert q.rules.pravidla[0].obrazek == rel_img
    assert q.image_note(rel_img) == "bod 1005 chybí"
    assert q.drawing_path_for_loading().name == "ukazkovy_vykres.dxf"


def test_stavy_chyb_a_porovnani(tmp_path, make_dxf):
    d = make_dxf(lambda msp, doc: (msp.add_line((0, 0), (10, 0)), msp.add_line((20, 0), (30, 0))))
    issues = check(d, "visici_konce")
    assert len(issues) == 4
    issues[0].state = "ignorovat"
    p = Project.create(tmp_path / "p")
    p.store_issues(issues)
    p.save()
    q = Project.open(tmp_path / "p")
    assert q.issue_states()[issues[0].key]["stav"] == "ignorovat"
    assert len(q.load_issues()) == 4
    d2 = make_dxf(lambda msp, doc: msp.add_line((0, 0), (10, 0)))
    new = check(d2, "visici_konce")
    c = compare(issues, new)
    assert (c.fixed, c.new, c.remaining) == (2, 0, 2)
    # dvě shodné chyby (stejné místo i text) se počítají dvakrát
    dup = issues[:1] * 2
    c = compare(dup, issues[:1])
    assert (c.fixed, c.new, c.remaining) == (1, 0, 1)


def test_vzor_pravidla_a_porovnani():
    from kontrola.io import read_dxf
    vzor = tpl.analyze(read_dxf(UKAZKY / "vzorovy_vykres.dxf"))
    kontrolovany = tpl.analyze(read_dxf(UKAZKY / "ukazkovy_vykres.dxf"))
    props = {r.kod: r for r in tpl.propose_rules(vzor)}
    # červená budova ze vzoru → číslo barvy MicroStationu 3
    assert props["BUDOVY"].geometrie == GeomType.POLYGON and props["BUDOVY"].barva == 3
    assert props["BUDOVY"].text.hladina == "POPIS_BUDOV"
    assert props["STROM_L"].blok == "STROM_L" and props["STROM_L"].povinne_atributy == ["DRUH"]
    assert props["PLOTY"].styl_cary == "DASHDOT"
    diffs = tpl.compare(vzor, kontrolovany)
    assert any(x.kategorie == "Hladina" and x.polozka == "POKUS" for x in diffs)
    assert any(x.kategorie == "Barva" and x.polozka == "BUDOVY" for x in diffs)


def test_exporty(tmp_path):
    from kontrola.export.dxf_export import LAYER, export_dxf
    from kontrola.export.pdf_report import export_pdf
    from kontrola.export.tables import export_csv, export_xlsx
    from kontrola.io import read_dxf

    d = read_dxf(UKAZKY / "ukazkovy_vykres.dxf")
    res = run_checks(d, RuleSet.load(UKAZKY / "konfigurace.yaml"), Config.load(UKAZKY / "konfigurace.yaml"))
    assert len(res.issues) > 10
    export_csv(res.issues, tmp_path / "a.csv")
    text = (tmp_path / "a.csv").read_text(encoding="utf-8-sig")
    assert text.startswith("Číslo;Typ kontroly;Závažnost")
    assert "Nezavřený polygon, mezera 0,03 m" in text
    export_xlsx(res.issues, tmp_path / "a.xlsx", "vykres.dxf")
    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "a.xlsx")
    assert wb.sheetnames == ["Chyby", "Souhrn"]
    assert wb["Chyby"].max_row == len(res.issues) + 1
    export_dxf(d, res.issues, tmp_path / "k.dxf")
    doc = ezdxf.readfile(tmp_path / "k.dxf")
    circles = doc.modelspace().query(f'CIRCLE[layer=="{LAYER}"]')
    assert len(circles) == len(res.issues)
    export_pdf(res.issues, tmp_path / "p.pdf", drawing_name="vykres.dxf")
    assert (tmp_path / "p.pdf").read_bytes()[:4] == b"%PDF"
    import pdfplumber
    with pdfplumber.open(tmp_path / "p.pdf") as pdf:
        first = pdf.pages[0].extract_text()
    assert "Protokol kontroly výkresu" in first and "Souhrn" in first


@pytest.mark.parametrize("value,expected", [("plocha", GeomType.POLYGON), ("Bodový", GeomType.BOD),
                                            ("linie", GeomType.LINIE), ("popis", GeomType.TEXT), ("?", None)])
def test_parsovani_typu_geometrie(value, expected):
    assert GeomType.parse(value) == expected


def test_detekce_hlavicky_bez_hlavicky():
    assert detect_header_row([["1", "2"], ["3", "4"]]) is None
    assert guess_mapping(["Kód prvku", "Hladina", "Barva"]) == {"kod": 0, "hladina": 1, "barva": 2}
