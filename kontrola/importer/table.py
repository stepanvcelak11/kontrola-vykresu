"""Import tabulky atributů od učitele (.xlsx, .csv, .pdf) a generování pravidel.

Postup:
1. :func:`read_table` načte surové řádky (u xlsx lze vybrat list),
2. :func:`detect_header_row` a :func:`guess_mapping` odhadnou řádek s hlavičkou
   a přiřazení sloupců podle názvů v hlavičce,
3. :func:`generate_rules` z řádků vytvoří pravidla a vrátí i seznam řádků,
   které se nepodařilo zpracovat.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from ..model import GeomType
from ..rules import (Rule, RuleSet, TextRule, geometry_from_element_types, normalize_linetype,
                     parse_allowed_values, parse_color, parse_element_types, split_list)

FIELDS: list[tuple[str, str]] = [
    ("kod", "Kód prvku"),
    ("nazev", "Název"),
    ("hladina", "Hladina"),
    ("barva", "Barva"),
    ("styl_cary", "Styl čáry"),
    ("tloustka", "Tloušťka čáry"),
    ("geometrie", "Typ geometrie"),
    ("povinne_atributy", "Povinné atributy"),
    ("povolene_hodnoty", "Povolené hodnoty"),
    ("blok", "Buňka / blok"),
    ("text_hladina", "Popis (text) na hladině"),
    ("styl_uzivatelsky", "Uživatelský styl čáry"),
    ("typy_prvku", "Typy prvků (MicroStation)"),
    ("font", "Font textu"),
    ("vyska_textu", "Výška textu"),
    ("sirka_textu", "Šířka textu"),
    ("zarovnani", "Zarovnání textu"),
    ("poznamka", "Poznámka"),
]
FIELD_LABELS = dict(FIELDS)

# klíčová slova v hlavičce (bez diakritiky, malými písmeny)
_KEYWORDS: dict[str, list[str]] = {
    "kod": ["kod", "code", "cislo prvku", "c. prvku", "id prvku", "znacka kodu"],
    "nazev": ["nazev", "popis prvku", "prvek", "objekt", "name", "vyznam", "jev", "trida prvku", "trida"],
    "hladina": ["hladina", "level", "vrstva", "layer", "lv", "vr", "cislo vrstvy"],
    "barva": ["barva", "color", "colour", "co", "ba"],
    "styl_cary": ["styl", "typ cary", "linetype", "line style", "druh cary", "lc", "st"],
    "tloustka": ["tloustka", "weight", "lineweight", "wt", "tl"],
    "geometrie": ["typ geometrie", "geometrie", "typ prvku", "typ objektu", "geometry", "druh", "typ"],
    "styl_uzivatelsky": ["us", "uzivatelsky styl", "uziv. styl", "custom style"],
    "typy_prvku": ["prvky", "typy prvku", "typ prvku ms", "element type", "typy kresebnych prvku"],
    "font": ["font", "pismo"],
    "vyska_textu": ["vyska", "vyska textu", "text height"],
    "sirka_textu": ["sirka", "sirka textu", "text width"],
    "zarovnani": ["zarovnani", "vztazny bod", "justification"],
    "povinne_atributy": ["povinne atributy", "atributy", "povinne", "attributes", "atribut"],
    "povolene_hodnoty": ["povolene hodnoty", "hodnoty", "ciselnik", "domena", "values", "povolene"],
    "blok": ["bunka", "blok", "cell", "block", "znacka", "symbol"],
    "text_hladina": ["popis na hladine", "hladina popisu", "text na hladine", "popis", "text"],
    "poznamka": ["poznamka", "note", "komentar"],
}


def strip_accents(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


def norm(s) -> str:
    return re.sub(r"\s+", " ", strip_accents(str(s or "")).lower()).strip()


@dataclass
class TableData:
    path: str
    sheets: list[str] = field(default_factory=list)
    sheet: str | None = None
    rows: list[list[str]] = field(default_factory=list)

    @property
    def ncols(self) -> int:
        return max((len(r) for r in self.rows), default=0)


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_table(path: str | Path, sheet: str | None = None) -> TableData:
    p = Path(path)
    suf = p.suffix.lower()
    if suf in (".xlsx", ".xlsm"):
        return _read_xlsx(p, sheet)
    if suf in (".csv", ".txt", ".tsv"):
        return _read_csv(p)
    if suf == ".pdf":
        return _read_pdf(p)
    if suf == ".xls":
        return _read_xls(p, sheet)
    raise ValueError(f"Nepodporovaný formát tabulky: {p.suffix}")


def _read_xlsx(p: Path, sheet: str | None) -> TableData:
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True, data_only=True)
    names = wb.sheetnames
    ws = wb[sheet] if sheet in names else wb[names[0]]
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append([_cell(v) for v in r])
    wb.close()
    while rows and not any(rows[-1]):
        rows.pop()
    return TableData(str(p), names, ws.title, _trim(rows))


def _read_xls(p: Path, sheet: str | None) -> TableData:
    """Starý formát Excelu (.xls, např. „Směrnice-výběr.xls“)."""
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover
        raise ValueError("Pro čtení .xls chybí knihovna xlrd – uložte tabulku jako .xlsx.") from exc
    wb = xlrd.open_workbook(str(p))
    names = wb.sheet_names()
    sh = wb.sheet_by_name(sheet) if sheet in names else wb.sheet_by_index(0)
    rows = [[_cell(c.value) for c in sh.row(r)] for r in range(sh.nrows)]
    while rows and not any(rows[-1]):
        rows.pop()
    return TableData(str(p), names, sh.name, _trim(rows))


def _read_csv(p: Path) -> TableData:
    raw = p.read_bytes()
    text = None
    for enc in ("utf-8-sig", "cp1250", "iso-8859-2"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw.decode("utf-8", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    return TableData(str(p), [], None, _trim(rows))


def _read_pdf(p: Path) -> TableData:
    import pdfplumber
    rows: list[list[str]] = []
    with pdfplumber.open(p) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables()
            if tables:
                for t in tables:
                    for r in t:
                        cells = [_cell(c).replace("\n", " ") for c in r]
                        if rows and _is_header_repeat(cells, rows):
                            continue
                        rows.append(cells)
            else:
                text = page.extract_text() or ""
                for line in text.splitlines():
                    parts = [x.strip() for x in re.split(r"\s{2,}|\t", line) if x.strip()]
                    if parts:
                        rows.append(parts)
    if not rows:
        raise ValueError("V PDF se nepodařilo najít žádnou tabulku ani text (může jít o sken – "
                         "pak tabulku přepište do Excelu).")
    return TableData(str(p), [], None, _trim(rows))


def _is_header_repeat(cells: list[str], rows: list[list[str]]) -> bool:
    """Hlavička opakovaná na každé stránce PDF."""
    hi = detect_header_row(rows)
    return hi is not None and [norm(c) for c in cells] == [norm(c) for c in rows[hi]]


def _trim(rows: list[list[str]]) -> list[list[str]]:
    n = max((len(r) for r in rows), default=0)
    rows = [r + [""] * (n - len(r)) for r in rows]
    # odstranit prázdné sloupce vpravo
    while n > 0 and all(not r[n - 1] for r in rows):
        rows = [r[:-1] for r in rows]
        n -= 1
    return rows


def _score_header_cell(cell: str) -> tuple[str | None, int]:
    c = norm(cell)
    if not c:
        return None, 0
    best, score = None, 0
    for fld, kws in _KEYWORDS.items():
        for k in kws:
            if c == k:
                s = 100 + len(k)
            elif len(k) > 2 and (c.startswith(k) or f" {k}" in f" {c}"):
                s = 10 + len(k)
            else:
                continue
            if s > score:
                best, score = fld, s
    return best, score


def detect_header_row(rows: list[list[str]], max_scan: int = 25) -> int | None:
    best, best_n = None, 0
    for i, r in enumerate(rows[:max_scan]):
        n = sum(1 for c in r if _score_header_cell(c)[0])
        if n > best_n:
            best, best_n = i, n
    return best if best_n >= 2 else None


def guess_mapping(header: list[str]) -> dict[str, int]:
    """Přiřadí sloupce polím podle hlavičky. Každý sloupec se použije jen jednou."""
    cands = []
    for col, cell in enumerate(header):
        c = norm(cell)
        if not c:
            continue
        for fld, kws in _KEYWORDS.items():
            for k in kws:
                if c == k:
                    s = 100 + len(k)
                elif len(k) > 2 and (c.startswith(k) or f" {k}" in f" {c}"):
                    s = 10 + len(k)
                else:
                    continue
                cands.append((s, fld, col))
    cands.sort(reverse=True)
    mapping: dict[str, int] = {}
    used: set[int] = set()
    for s, fld, col in cands:
        if fld in mapping or col in used:
            continue
        mapping[fld] = col
        used.add(col)
    return mapping


@dataclass
class ImportResult:
    rules: RuleSet
    errors: list[str] = field(default_factory=list)  # řádky, které nešlo zpracovat
    warnings: list[str] = field(default_factory=list)  # zpracováno, ale s výhradou
    processed: int = 0

    def summary(self) -> str:
        s = f"Vzniklo {len(self.rules.pravidla)} pravidel."
        if self.errors:
            s += f" Nepodařilo se zpracovat {len(self.errors)} řádků."
        if self.warnings:
            s += f" Upozornění: {len(self.warnings)}."
        return s


def _num(txt: str) -> float | None:
    m = re.search(r"-?\d+(?:[.,]\d+)?", txt or "")
    return float(m.group(0).replace(",", ".")) if m else None


def generate_rules(rows: list[list[str]], header_row: int | None, mapping: dict[str, int],
                   source: str = "") -> ImportResult:
    rs = RuleSet()
    res = ImportResult(rs)
    seen: dict[str, int] = {}
    start = (header_row + 1) if header_row is not None else 0
    if "kod" not in mapping and "hladina" not in mapping and "blok" not in mapping:
        res.errors.append("Není přiřazen sloupec s kódem, hladinou ani buňkou – pravidla nelze vytvořit.")
        return res

    def get(r: list[str], fld: str) -> str:
        col = mapping.get(fld)
        if col is None or col >= len(r):
            return ""
        return (r[col] or "").strip()

    for i in range(start, len(rows)):
        r = rows[i]
        rowno = i + 1
        filled = [c for c in r if str(c).strip()]
        if len(filled) <= 1:
            continue  # prázdný řádek nebo nadpis skupiny
        kod = get(r, "kod")
        nazev = get(r, "nazev")
        hladina = get(r, "hladina")
        blok = get(r, "blok")
        if not kod:
            if "kod" in mapping:
                res.errors.append(f"Řádek {rowno}: chybí kód prvku ({nazev or ' / '.join(filled[:3])}).")
                continue
            kod = nazev or hladina or blok  # bez sloupce s kódem je kódem název prvku
        if not re.search(r"[0-9A-Za-zÀ-ž]", kod):
            res.errors.append(f"Řádek {rowno}: neplatný kód „{kod}“ ({nazev}).")
            continue
        if not hladina and not blok and "kod" not in mapping:
            res.errors.append(f"Řádek {rowno}: chybí hladina i buňka.")
            continue
        if kod in seen:
            res.errors.append(f"Řádek {rowno}: kód {kod} už je na řádku {seen[kod]} – přeskočeno.")
            continue
        warn = []
        geom_txt = get(r, "geometrie")
        geom = GeomType.parse(geom_txt)
        types = parse_element_types(get(r, "typy_prvku"))
        if get(r, "typy_prvku") and not types:
            warn.append(f"nerozpoznané typy prvků „{get(r, 'typy_prvku')}“")
        if geom is None and types:
            geom = geometry_from_element_types(types)
        if geom_txt and geom is None:
            warn.append(f"nerozpoznaný typ geometrie „{geom_txt}“")
        if geom is None and blok:
            geom = GeomType.BOD
        barva_txt = get(r, "barva")
        barva = parse_color(barva_txt)
        if barva_txt and barva is None:
            warn.append(f"nerozpoznaná barva „{barva_txt}“")
        tl_txt = get(r, "tloustka").replace(",", ".")
        tl = None
        if tl_txt:
            m = re.search(r"\d+(\.\d+)?", tl_txt)
            if m:
                tl = float(m.group(0))
            else:
                warn.append(f"nerozpoznaná tloušťka „{tl_txt}“")
        attrs = [a.upper() for a in split_list(get(r, "povinne_atributy"))]
        allowed = parse_allowed_values(get(r, "povolene_hodnoty"), attrs[0] if attrs else None)
        if get(r, "povolene_hodnoty") and not allowed:
            warn.append("povolené hodnoty nemají přiřazený atribut (vyplňte povinný atribut nebo "
                        "zapište ATRIBUT=hodnota1,hodnota2)")
        text_layer = get(r, "text_hladina")
        text_rule = None
        if text_layer and geom != GeomType.TEXT:
            text_rule = TextRule(povinny=True, hladina=text_layer)
        # uživatelský styl čáry (např. „2.123“ z ugeo_vp.rsc) má přednost před základním stylem
        styl = get(r, "styl_uzivatelsky") or get(r, "styl_cary")
        font = get(r, "font")
        font = re.sub(r"^\s*\d+\s*[-–]\s*", "", font) if font else ""  # „1 - CS_WORKING“ → „CS_WORKING“
        rule = Rule(kod=kod, nazev=nazev, geometrie=geom, hladina=hladina or None, barva=barva,
                    styl_cary=normalize_linetype(styl) if styl and "," not in styl else (styl or None),
                    tloustka=tl, blok=blok or None,
                    povinne_atributy=attrs, povolene_hodnoty=allowed, text=text_rule,
                    typy_prvku=types, vyska_textu=_num(get(r, "vyska_textu")),
                    sirka_textu=_num(get(r, "sirka_textu")), font=font or None,
                    zarovnani=get(r, "zarovnani") or None,
                    topologie=not re.search(r"\bvstup", nazev, re.IGNORECASE),
                    poznamka=get(r, "poznamka") or None,
                    zdroj=f"{Path(source).name}, řádek {rowno}" if source else f"řádek {rowno}")
        rs.pravidla.append(rule)
        seen[kod] = rowno
        res.processed += 1
        if warn:
            res.warnings.append(f"Řádek {rowno} ({kod}): " + "; ".join(warn) + ".")
    # hladiny popisů patří mezi povolené
    for r in rs.pravidla:
        if r.text and r.text.hladina and not any(
                x.hladina and x.hladina.upper() == r.text.hladina.upper() for x in rs.pravidla):
            if r.text.hladina not in rs.povolene_hladiny:
                rs.povolene_hladiny.append(r.text.hladina)
    return res


def import_table(path: str | Path, sheet: str | None = None) -> tuple[TableData, int | None, dict[str, int]]:
    """Načte tabulku a odhadne hlavičku i přiřazení sloupců."""
    td = read_table(path, sheet)
    hi = detect_header_row(td.rows)
    mapping = guess_mapping(td.rows[hi]) if hi is not None else {}
    return td, hi, mapping
