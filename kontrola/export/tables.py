"""Export seznamu chyb do CSV a Excelu."""

from __future__ import annotations

import csv
import datetime as dt
from collections import Counter
from pathlib import Path

from ..checks.base import Issue, Severity

from ..navody import navod

HEADER = ["Číslo", "Typ kontroly", "Závažnost", "Popis", "Vrstva", "X", "Y", "Stav", "Poznámka", "Handle",
          "Jak opravit"]


def _row(i: Issue) -> list:
    return [i.number, i.check_name, i.severity.value, i.message, i.layer, round(i.x, 3), round(i.y, 3),
            i.state, i.note, ",".join(i.handles), navod(i)]


def export_csv(issues: list[Issue], path: str | Path):
    """CSV se středníkem a desetinnou čárkou – otevře se správně v českém Excelu."""
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(HEADER)
        for i in issues:
            r = _row(i)
            r[5] = f"{i.x:.3f}".replace(".", ",")
            r[6] = f"{i.y:.3f}".replace(".", ",")
            w.writerow(r)


def export_xlsx(issues: list[Issue], path: str | Path, drawing_name: str = ""):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Chyby"
    ws.append(HEADER)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="44546A")
        c.alignment = Alignment(vertical="center")
    fills = {Severity.CHYBA: "F8D7DA", Severity.VAROVANI: "FFE8CC", Severity.INFO: "DCE9FA"}
    for i in issues:
        ws.append(_row(i))
        ws.cell(ws.max_row, 3).fill = PatternFill("solid", fgColor=fills[i.severity])
        for col in (6, 7):
            ws.cell(ws.max_row, col).number_format = "0.000"
    widths = [8, 26, 11, 60, 18, 14, 14, 11, 30, 14, 90]
    for k, wdt in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(k)].width = wdt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}{max(1, ws.max_row)}"

    s = wb.create_sheet("Souhrn")
    s.append(["Kontrola výkresu – souhrn"])
    s["A1"].font = Font(bold=True, size=14)
    s.append(["Výkres", drawing_name])
    s.append(["Datum", dt.datetime.now().strftime("%d.%m.%Y %H:%M")])
    s.append([])
    s.append(["Typ kontroly", "Chyby", "Varování", "Info", "Celkem"])
    for c in s[5]:
        c.font = Font(bold=True)
    per: dict[str, Counter] = {}
    for i in issues:
        per.setdefault(i.check_name, Counter())[i.severity] += 1
    for name in sorted(per):
        c = per[name]
        s.append([name, c[Severity.CHYBA], c[Severity.VAROVANI], c[Severity.INFO], sum(c.values())])
    tot = Counter(i.severity for i in issues)
    s.append(["Celkem", tot[Severity.CHYBA], tot[Severity.VAROVANI], tot[Severity.INFO], len(issues)])
    for c in s[s.max_row]:
        c.font = Font(bold=True)
    s.column_dimensions["A"].width = 32
    s.column_dimensions["B"].width = 40
    wb.save(path)
